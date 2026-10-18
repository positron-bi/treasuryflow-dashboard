from __future__ import annotations

import os
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import time
import unittest
from unittest.mock import patch

from dashboard_auth import protect_html
from pipeline.cash_accounts import read_cash_accounts
from pipeline.make_html_v2 import build_html
from openpyxl import Workbook

from treasuryflow_core import (
    SourceSet,
    find_sources,
    normalize_date,
    source_signature,
    clear_upload_request,
    has_upload_request,
    mark_upload_request,
)


class TreasuryFlowCoreTests(unittest.TestCase):
    def test_cash_accounts_follow_headers_after_column_insert(self) -> None:
        sheet = Workbook().active
        for index, value in enumerate(["Co.", "Bank - Account", "Account Number", "Bank Balance", "Deposit", "Inserted", "Bank Available", "Deposit"], 1):
            sheet.cell(5, index, value)
        for index, value in enumerate(["Pars Anahid", "Bank A", "123", 900, 700, 999999, 200, 333], 1):
            sheet.cell(6, index, value)

        accounts = read_cash_accounts(sheet, {"Pars Anahid"}, str, 10)
        self.assertEqual(accounts, [{"co": "Pars Anahid", "bank": "Bank A", "acct": "123", "bal": 20.0, "blocked": 70.0}])

    def test_upload_request_lifecycle(self) -> None:
        with TemporaryDirectory() as temp:
            home = Path(temp)
            marker = mark_upload_request(home, [home / "daily.xlsx"])
            self.assertTrue(marker.exists())
            self.assertTrue(has_upload_request(home))
            self.assertIn("daily.xlsx", marker.read_text(encoding="utf-8"))
            clear_upload_request(home)
            self.assertFalse(has_upload_request(home))

    def test_html_builder_adds_generation_timestamp(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            data = root / "data.json"
            output = root / "dashboard.html"
            data.write_text(json.dumps({"report_date": "1405/05/28"}), encoding="utf-8")
            build_html(data, output)
            html = output.read_text(encoding="utf-8")
            self.assertIn('"generated_at":', html)
            self.assertIn('id="updatedBadge"', html)
            self.assertIn('id="claudeAnalyzeBtn"', html)
            self.assertIn('function openClaudeAnalysis()', html)
            self.assertIn('https://claude.ai/new', html)
            self.assertIn('TreasuryFlow_Claude_', html)
            self.assertIn('id="treasuryChatBtn"', html)
            self.assertIn('function context(question)', html)
            self.assertIn('https://api.groq.com/openai/v1/chat/completions', html)

    def test_protected_html_hides_plain_content_and_password(self) -> None:
        with patch("dashboard_auth.load_users", return_value=[{"username": "TestUser", "password": "Secret123"}]):
            protected = protect_html("<h1>financial dashboard</h1>", "unused.xlsx")
        self.assertIn("const CFG=", protected)
        self.assertNotIn("financial dashboard", protected)
        self.assertNotIn("Secret123", protected)

    def test_normalize_solar_date(self) -> None:
        self.assertEqual(normalize_date("گزارش ۱۴۰۵-۵-۲۸"), "1405/05/28")
        self.assertEqual(normalize_date("1405/05/28"), "1405/05/28")
        self.assertIsNone(normalize_date("1405/13/01"))

    def test_find_sources_selects_latest_report_date(self) -> None:
        with TemporaryDirectory() as temp:
            home = Path(temp)
            older = home / "Treasury Daily Report - 1405-05-28.xlsx"
            newer = home / "Treasury Daily Report - 1405-06-01.xlsx"
            older.write_bytes(b"older")
            newer.write_bytes(b"newer")
            # حتی اگر mtime فایل قدیمی جدیدتر باشد، تاریخ گزارش معیار اصلی است.
            now = time.time()
            os.utime(older, (now + 10, now + 10))
            (home / "Manual_Inputs.xlsx").write_bytes(b"manual")
            (home / "گزارش تسهیلات.xlsx").write_bytes(b"facilities")

            sources = find_sources(home)
            self.assertEqual(sources.daily, newer)
            self.assertEqual(sources.manual.name, "Manual_Inputs.xlsx")
            self.assertEqual(sources.facilities.name, "گزارش تسهیلات.xlsx")

    def test_signature_changes_when_content_changes(self) -> None:
        with TemporaryDirectory() as temp:
            daily = Path(temp) / "daily.xlsx"
            daily.write_bytes(b"v1")
            sources = SourceSet(daily=daily, manual=None, facilities=None)
            first = source_signature(sources)
            daily.write_bytes(b"v2")
            second = source_signature(sources)
            self.assertNotEqual(first, second)


if __name__ == "__main__":
    unittest.main()
